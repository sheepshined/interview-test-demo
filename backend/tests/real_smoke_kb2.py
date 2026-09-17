"""v0.8 知识库增强真实冒烟 (手动): 上传/网址/检索/语义关联/覆盖检测全链路。

前置: 后端已启动最新代码; LLM_API_KEY 已配 (结构化/分类走 STRONG 模型)。
真实 BGE 嵌入 + 真实 Chroma kb_notes 集合 (开发库, 结束后清理)。
"""
import io
import os
import sys

import httpx

BASE = f"http://127.0.0.1:{os.getenv('SMOKE_PORT', '8000')}"


def make_pptx() -> bytes:
    from pptx import Presentation
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Transformer 自注意力"
    s.placeholders[1].text = "Q K V 三矩阵投影\n缩放点积注意力 Scaled Dot-Product"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def main():
    http = httpx.Client(base_url=BASE, timeout=300)
    uname = "kb8_smoke_user"
    r = http.post("/api/login", json={"username": uname, "password": "kb8pass123"})
    if r.status_code != 200:
        r = http.post("/api/register", json={"username": uname, "password": "kb8pass123"})
        assert r.status_code == 200, r.text
        r = http.post("/api/login", json={"username": uname, "password": "kb8pass123"})
    h = {"Authorization": f"Bearer {r.json()['token']}"}
    print("[1] login OK")

    # 清旧数据
    for n in http.get("/api/kb/notes", headers=h).json().get("notes", []):
        http.delete(f"/api/kb/notes/{n['id']}", headers=h)

    # 2. 上传 pptx + md (LLM 结构化; Key 余额不足时自动降级原文直入)
    r = http.post("/api/kb/notes/upload", headers=h, files=[
        ("files", ("attention.pptx", make_pptx(), "application/vnd.openxmlformats-officedocument.presentationml.presentation")),
        ("files", ("rag_notes.md", "# RAG 复习\nRAG 检索增强生成结合向量检索召回知识再由大模型生成答案".encode(), "text/markdown")),
    ])
    assert r.status_code == 200, r.text
    results = r.json()["results"]
    llm_ok = True
    for e in results:
        print(f"    upload {e['file']}: {e['status']} llm={e.get('llm_used')} title={e.get('title')!r}")
        assert e["status"] == "created", e
        llm_ok = llm_ok and bool(e.get("llm_used"))
    if llm_ok:
        print("[2] upload with real LLM structuring OK")
    else:
        print("[2] upload OK (LLM 降级原文直入 — 检查 LLM_API_KEY 余额/网络)")

    # 3. 网址收藏
    r = http.post("/api/kb/notes/url", headers=h, json={
        "url": "https://arxiv.org/abs/1706.03762",
        "title": "Attention Is All You Need",
        "description": "Transformer 原始论文, 提出自注意力机制取代 RNN",
        "category": "深度学习",
    })
    assert r.status_code == 200, r.text
    print("[3] url note OK")

    # 4. 向量检索 (真实 BGE)
    r = http.get("/api/kb/search", params={"q": "自注意力机制 QKV"}, headers=h)
    data = r.json()
    assert data["success"] and not data["sparse"], data
    print(f"[4] search OK: top={data['results'][0]['title']!r} score={data['results'][0]['score']}")
    r = http.get("/api/kb/search", params={"q": "怎么做红烧肉"}, headers=h)
    assert r.json()["sparse"] is True
    print("[5] sparse query correctly flagged OK")

    # 5. 重建语义关联 (真实向量两两相似度)
    r = http.post("/api/kb/rebuild", headers=h)
    stats = r.json()
    print(f"[6] rebuild OK: {stats['notes']} notes, {stats['sem_links']} sem_links")

    # 6. 图谱: 语义边 + 类型
    g = http.get("/api/kb/graph", headers=h).json()["graph"]
    sem = [e for e in g["edges"] if e["kind"] == "semantic"]
    types = {n["note_type"] for n in g["nodes"]}
    print(f"[7] graph OK: {len(g['nodes'])} nodes, types={types}, {len(sem)} semantic edges")
    assert "url" in types and "file" in types

    # 7. from-report 覆盖检测
    r = http.post("/api/kb/notes/from-report", headers=h, json={
        "items": [
            {"topic": "Transformer 注意力机制", "suggestion": "复习 QKV 计算"},
            {"topic": "红烧肉烹饪", "suggestion": "学习东坡肉技法"},
        ], "report_id": "smoke_rpt",
    })
    cov = r.json()["coverage"]
    print("    raw coverage:", cov)
    by_topic = {c["topic"]: c for c in cov}
    assert not by_topic["Transformer 注意力机制"]["sparse"]       # 知识库已覆盖
    assert by_topic["红烧肉烹饪"]["sparse"]                        # 缺失提示
    print(f"[8] coverage OK: covered={len(by_topic['Transformer 注意力机制']['matched'])}, sparse flagged")

    # 清理
    for n in http.get("/api/kb/notes", headers=h).json().get("notes", []):
        http.delete(f"/api/kb/notes/{n['id']}", headers=h)
    print("[9] cleanup OK\n\nAll v0.8 smoke checks passed.")


if __name__ == "__main__":
    try:
        main()
    except AssertionError as exc:
        print(f"FAILED: {exc}")
        sys.exit(1)
