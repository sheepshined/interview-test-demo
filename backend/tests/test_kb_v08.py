"""v0.8 知识库增强测试: 文件提取 / 新字段 / 网址 / 分类 / 图谱边分类 / 向量写入钩子。

LLM 调用 (structure_file_note/suggest_category/tidy) 全部 mock — 单测不依赖外部 Key。
运行环境由 conftest 隔离 (KB_DB_PATH / KB_CHROMA_PATH 指向临时目录)。
"""
import io
from unittest.mock import patch

import pytest
from fastapi.testclient import TestClient

import kb_extract
import knowledge

from server import app

client = TestClient(app)


class FakeEmbeddings:
    """确定性伪嵌入: 关键词重叠→向量接近。维度固定 16, 避免真实 BGE 768 维锁死集合。"""
    WORDS = ["rag", "检索", "向量", "mysql", "索引", "算法", "排序", "redis"]

    def embed_query(self, text: str):
        vec = [0.0] * 16
        low = (text or "").lower()
        for word in self.WORDS:
            if word in low:
                vec[abs(hash(word)) % 16] += 1.0
        norm = sum(v * v for v in vec) ** 0.5 or 1.0
        return [v / norm for v in vec]


@pytest.fixture(autouse=True)
def fake_embeddings():
    """全文件统一伪嵌入: 不加载真实 BGE (速度), 且向量维度一致 (Chroma 集合维度锁定)。"""
    with patch("kb_vectors._embedding_cache", [FakeEmbeddings()]):
        yield


@pytest.fixture(autouse=True)
def headers():
    import itertools
    counter = getattr(headers, "_counter", itertools.count())
    headers._counter = counter
    username = f"kb8user{next(counter)}"
    resp = client.post("/api/register", json={"username": username, "password": "pass12345"})
    assert resp.status_code == 200
    return {"Authorization": f"Bearer {resp.json()['token']}"}


# ============================================================
# kb_extract: 文本 / pptx / docx
# ============================================================

def test_extract_plain_text_utf8_and_gbk():
    r = kb_extract.extract_file("note.md", "# 你好".encode("utf-8"))
    assert r["extract_mode"] == "text" and "你好" in r["raw_text"]
    r = kb_extract.extract_file("old.txt", "你好".encode("gbk"))
    assert "你好" in r["raw_text"]


def test_extract_rejects_unsupported():
    with pytest.raises(ValueError, match="不支持"):
        kb_extract.extract_file("evil.exe", b"bin")


def _make_pptx() -> bytes:
    from pptx import Presentation
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])   # 标题+内容
    slide.shapes.title.text = "RAG 入门"
    slide.placeholders[1].text = "检索增强生成\n结合 [[向量检索]]"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _make_docx() -> bytes:
    import docx
    doc = docx.Document()
    doc.add_heading("索引原理", level=1)
    doc.add_paragraph("B+ 树与最左前缀")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_extract_pptx_and_docx():
    r = kb_extract.extract_file("rag.pptx", _make_pptx())
    assert r["extract_mode"] == "pptx" and r["pages"] == 1
    assert "RAG 入门" in r["raw_text"] and "向量检索" in r["raw_text"]

    r = kb_extract.extract_file("index.docx", _make_docx())
    assert r["extract_mode"] == "docx"
    assert "索引原理" in r["raw_text"] and "B+ 树" in r["raw_text"]


# ============================================================
# 上传端点 (LLM mock)
# ============================================================

def _mock_structure(file_name, raw_text):
    return {"title": f"AI:{raw_text[:10]}", "markdown": f"## 整理\n{raw_text}",
            "category": "测试分类", "tags": ["t1"], "llm_used": True}


def test_upload_md_with_llm_structuring(headers):
    with patch("kb_llm.structure_file_note", side_effect=_mock_structure):
        resp = client.post(
            "/api/kb/notes/upload",
            files={"files": ("quickstart.md", "# RAG 指南\n核心内容".encode(), "text/markdown")},
            headers=headers,
        )
    assert resp.status_code == 200
    results = resp.json()["results"]
    assert results[0]["status"] == "created"
    assert results[0]["llm_used"] is True

    note = client.get(f"/api/kb/notes/{results[0]['note_id']}", headers=headers).json()["note"]
    assert note["note_type"] == "file"
    assert note["category"] == "测试分类"
    assert note["file_name"] == "quickstart.md"
    assert "整理" in note["content"]


def test_upload_llm_failure_degrades_to_raw(headers):
    # LLM 降级: 返回 llm_used=False 时标题=文件名, 内容=原文
    with patch("kb_llm.structure_file_note",
               return_value={"title": "fallback.md"[:-3], "markdown": "原文",
                             "category": "", "tags": [], "llm_used": False}):
        resp = client.post(
            "/api/kb/notes/upload",
            files={"files": ("fallback.md", b"raw content here", "text/markdown")},
            headers=headers,
        )
    assert resp.status_code == 200
    entry = resp.json()["results"][0]
    assert entry["status"] == "created" and entry["llm_used"] is False


def test_upload_unsupported_and_duplicate(headers):
    with patch("kb_llm.structure_file_note", side_effect=_mock_structure):
        # 不支持的格式 → failed
        resp = client.post(
            "/api/kb/notes/upload",
            files={"files": ("a.exe", b"bin", "application/octet-stream")},
            headers=headers,
        )
        assert resp.json()["results"][0]["status"] == "failed"

        # 首次上传 → created
        resp = client.post(
            "/api/kb/notes/upload",
            files={"files": ("dup.md", b"content", "text/markdown")},
            headers=headers,
        )
        assert resp.json()["results"][0]["status"] == "created"

        # 同名重复上传 → skipped
        resp = client.post(
            "/api/kb/notes/upload",
            files={"files": ("dup.md", b"content", "text/markdown")},
            headers=headers,
        )
        assert resp.json()["results"][0]["status"] == "skipped"


# ============================================================
# 网址笔记
# ============================================================

def test_url_note_requires_description(headers):
    resp = client.post("/api/kb/notes/url", headers=headers,
                       json={"url": "https://example.com", "description": ""})
    assert resp.status_code == 400
    assert "描述" in resp.json()["message"]


def test_url_note_created_with_link_back(headers):
    resp = client.post("/api/kb/notes/url", headers=headers,
                       json={"url": "https://bge.example.com", "title": "BGE 模型卡",
                             "description": "中文嵌入模型, 余弦检索用",
                             "category": "大模型"})
    assert resp.status_code == 200
    note = client.get(f"/api/kb/notes/{resp.json()['id']}", headers=headers).json()["note"]
    assert note["note_type"] == "url"
    assert note["source_url"] == "https://bge.example.com"
    assert "https://bge.example.com" in note["content"]   # 内容带可点的来源链接
    assert note["category"] == "大模型"


# ============================================================
# 分类
# ============================================================

def test_categories_aggregation(headers):
    client.post("/api/kb/notes", headers=headers,
                json={"title": "A", "content": "x", "category": "算法"})
    client.post("/api/kb/notes", headers=headers,
                json={"title": "B", "content": "x", "category": "算法"})
    client.post("/api/kb/notes", headers=headers,
                json={"title": "C", "content": "x", "category": ""})
    resp = client.get("/api/kb/categories", headers=headers)
    cats = resp.json()["categories"]
    assert cats == [{"category": "算法", "count": 2}]   # 空分类不计


def test_list_filter_by_category(headers):
    client.post("/api/kb/notes", headers=headers,
                json={"title": "甲", "content": "x", "category": "前端"})
    client.post("/api/kb/notes", headers=headers,
                json={"title": "乙", "content": "x", "category": "后端"})
    resp = client.get("/api/kb/notes", params={"category": "前端"}, headers=headers)
    titles = [n["title"] for n in resp.json()["notes"]]
    assert titles == ["甲"]


# ============================================================
# 图谱: 类型/分类/边 kind
# ============================================================

def test_graph_node_types_and_category_filter(headers):
    client.post("/api/kb/notes", headers=headers, json={"title": "N1", "content": "链 [[N2]]", "category": "X"})
    client.post("/api/kb/notes/url", headers=headers,
                json={"url": "https://u.com", "title": "U1", "description": "d"})
    with patch("kb_llm.structure_file_note", side_effect=_mock_structure):
        client.post("/api/kb/notes/upload", headers=headers,
                    files={"files": ("f.md", b"file content", "text/markdown")})

    graph = client.get("/api/kb/graph", headers=headers).json()["graph"]
    types = {n["id"]: n["note_type"] for n in graph["nodes"]}
    assert types["N1"] == "note" and types["U1"] == "url"
    assert any(t == "file" for t in types.values())
    assert all(e["kind"] == "wiki" for e in graph["edges"])   # 语义边默认无数据

    # 分类过滤: 只看 X → 只剩 N1 (其出链 N2 为虚节点保留)
    graph = client.get("/api/kb/graph", params={"category": "X"}, headers=headers).json()["graph"]
    titles = {n["id"] for n in graph["nodes"]}
    assert titles == {"N1", "N2"}


def test_semantic_edges_rendered_from_table(headers):
    client.post("/api/kb/notes", headers=headers, json={"title": "S1", "content": "a"})
    client.post("/api/kb/notes", headers=headers, json={"title": "S2", "content": "b"})
    # 直接写 sem_links 模拟预计算结果
    with knowledge._connect() as conn:
        conn.execute(
            "INSERT INTO sem_links (username, title_a, title_b, score, created_at) VALUES (?,?,?,?,?)",
            (headers and _current_username(headers), "S1", "S2", 0.72, "now"),
        )
    graph = client.get("/api/kb/graph", headers=headers).json()["graph"]
    sem = [e for e in graph["edges"] if e["kind"] == "semantic"]
    assert len(sem) == 1 and sem[0]["score"] == 0.72

    # 关闭语义边
    graph = client.get("/api/kb/graph", params={"semantic": "false"}, headers=headers).json()["graph"]
    assert all(e["kind"] == "wiki" for e in graph["edges"])


def _current_username(h):
    token = h["Authorization"].split(" ")[1]
    import auth as auth_mod
    payload = auth_mod.verify_token(token)
    return payload["username"]


# ============================================================
# 向量检索钩子 (embedding 已由 autouse fixture 统一 mock)
# ============================================================

def test_search_and_sparse_flag(headers):
    client.post("/api/kb/notes", headers=headers,
                json={"title": "RAG 检索", "content": "向量 检索 rag"})
    resp = client.get("/api/kb/search", params={"q": "rag 检索"}, headers=headers)
    data = resp.json()
    assert data["success"] and not data["sparse"]
    assert data["results"][0]["title"] == "RAG 检索"

    # 完全无关查询 → sparse 提示导入
    resp = client.get("/api/kb/search", params={"q": "量子物理"}, headers=headers)
    assert resp.json()["sparse"] is True


def test_from_report_returns_coverage(headers):
    resp = client.post("/api/kb/notes/from-report", headers=headers, json={
        "items": [{"topic": "RAG 检索优化", "suggestion": "学分块"}],
        "report_id": "rpt_x",
    })
    data = resp.json()
    assert data["success"]
    assert len(data["coverage"]) == 1
    assert "topic" in data["coverage"][0] and "matched" in data["coverage"][0]
