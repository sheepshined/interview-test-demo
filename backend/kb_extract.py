"""
kb_extract.py — 知识库文件原文提取 (v0.8 批次A)

支持格式:
  .md/.txt/.markdown  → 直接读文本 (UTF-8, GBK 回退)
  .pptx               → python-pptx 逐页提取 标题+正文+表格
  .docx               → python-docx 段落+表格
  .pdf                → pdfplumber 文本层; 提取为空/过少 → 千问视觉转写 (复用 vision_parser 渲染)

统一产出 {"raw_text": str, "pages": int, "extract_mode": "text|pptx|docx|pdf|vision"}。
原始文件不落盘 (调用方传 bytes → 本模块写临时文件供解析器读取, 用完即删)。
"""
import logging
import os
import tempfile
from typing import Dict, List, Optional

import config

logger = logging.getLogger(__name__)

# 扫描 PDF 视觉兜底: 页数上限 (知识库文档比简历长, 放宽到 20 页)
_VISION_MAX_PAGES = 20
# PDF 文本层最少字符数: 少于该值判定为扫描件走视觉
_PDF_MIN_TEXT_CHARS = 50

SUPPORTED_EXTENSIONS = {".md", ".markdown", ".txt", ".pptx", ".docx", ".pdf"}


def extract_file(file_name: str, data: bytes) -> Dict:
    """按扩展名提取文件内容。

    Returns:
        {"raw_text": str, "pages": int, "extract_mode": str}
    Raises:
        ValueError: 不支持的格式 / 提取失败 (信息给用户看)
    """
    ext = os.path.splitext(file_name or "")[1].lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"不支持的格式 {ext}，支持: md/txt/pptx/docx/pdf")

    # 文本类: 不需要临时文件
    if ext in {".md", ".markdown", ".txt"}:
        return _extract_plain(data)

    # 二进制类: 写临时文件供各解析器读取
    tmp_path = os.path.join(tempfile.mkdtemp(prefix="kb_upload_"), f"file{ext}")
    with open(tmp_path, "wb") as f:
        f.write(data)
    try:
        if ext == ".pptx":
            return _extract_pptx(tmp_path)
        if ext == ".docx":
            return _extract_docx(tmp_path)
        if ext == ".pdf":
            return _extract_pdf(tmp_path, file_name)
        raise ValueError(f"不支持的格式 {ext}")
    finally:
        try:
            os.remove(tmp_path)
        except OSError:
            pass


# ============================================================
# 各格式提取
# ============================================================

def _extract_plain(data: bytes) -> Dict:
    for encoding in ("utf-8", "gbk"):
        try:
            text = data.decode(encoding)
            return {"raw_text": text, "pages": 1, "extract_mode": "text"}
        except UnicodeDecodeError:
            continue
    raise ValueError("无法识别文件编码 (支持 UTF-8 / GBK)")


def _extract_pptx(path: str) -> Dict:
    from pptx import Presentation

    try:
        prs = Presentation(path)
    except Exception as exc:
        raise ValueError(f"PPT 文件损坏或无法解析: {exc}") from exc

    blocks: List[str] = []
    for idx, slide in enumerate(prs.slides, 1):
        parts: List[str] = []
        # 文本框 (含标题占位符)
        for shape in slide.shapes:
            if shape.has_text_frame:
                lines = [
                    p.text.strip()
                    for p in shape.text_frame.paragraphs
                    if p.text and p.text.strip()
                ]
                if lines:
                    parts.append("\n".join(lines))
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))
        blocks.append(f"## 第{idx}页\n" + ("\n\n".join(parts) if parts else "(本页无文本)"))

    raw = "\n".join(blocks)
    if not raw.strip():
        raise ValueError("PPT 中未提取到任何文本 (可能全部为图片)")
    return {"raw_text": raw, "pages": len(blocks), "extract_mode": "pptx"}


def _extract_docx(path: str) -> Dict:
    import docx

    try:
        doc = docx.Document(path)
    except Exception as exc:
        raise ValueError(f"Word 文件损坏或无法解析: {exc}") from exc

    parts: List[str] = []
    for para in doc.paragraphs:
        if para.text and para.text.strip():
            # 保留标题层级
            if para.style and para.style.name and para.style.name.startswith("Heading"):
                level = para.style.name.replace("Heading ", "").strip()
                try:
                    parts.append("#" * max(1, min(4, int(level))) + " " + para.text.strip())
                except ValueError:
                    parts.append("## " + para.text.strip())
            else:
                parts.append(para.text.strip())
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    raw = "\n\n".join(parts)
    if not raw.strip():
        raise ValueError("Word 文档中未提取到任何文本")
    return {"raw_text": raw, "pages": max(1, len(parts) // 30), "extract_mode": "docx"}


def _extract_pdf(path: str, file_name: str) -> Dict:
    """PDF: 文本层优先, 扫描件走千问视觉转写。"""
    import pdfplumber

    texts: List[str] = []
    try:
        with pdfplumber.open(path) as pdf:
            for page in pdf.pages:
                texts.append(page.extract_text() or "")
    except Exception as exc:
        raise ValueError(f"PDF 无法解析: {exc}") from exc

    raw = "\n".join(texts).strip()
    if len(raw) >= _PDF_MIN_TEXT_CHARS:
        return {"raw_text": raw, "pages": len(texts), "extract_mode": "pdf"}

    # 文本层过少 → 视觉兜底 (效果优先: 用户已确认扫描件也走模型)
    vision = _extract_pdf_by_vision(path, file_name)
    if vision:
        return vision
    raise ValueError("PDF 无法提取文本（疑似扫描件），且视觉模型未配置或解析失败，请配置 DASHSCOPE_API_KEY 后重试")


def _extract_pdf_by_vision(path: str, file_name: str) -> Optional[Dict]:
    """千问视觉转写扫描件 PDF。复用 vision_parser 的渲染管线。"""
    try:
        from resume.vision_parser import pdf_to_base64_images, vision_available
    except ImportError:
        return None
    if not vision_available():
        logger.info("视觉模型不可用, 扫描件 PDF 无法兜底")
        return None

    images = pdf_to_base64_images(path, max_pages=_VISION_MAX_PAGES)
    if not images:
        return None

    system = (
        "你是文档转写助手。用户给出文档页面图片, 请将图中所有文字忠实转写为 Markdown "
        "(保留标题层级/列表/表格结构), 不得虚构不存在的内容。只输出转写结果。"
    )
    content: List[Dict] = [{"type": "text", "text": f"转写文档《{file_name}》的 {len(images)} 页:"}]
    for img in images:
        content.append({"type": "image_url", "image_url": {"url": f"data:image/png;base64,{img}"}})

    try:
        from openai import OpenAI
        client = OpenAI(api_key=config.DASHSCOPE_API_KEY, base_url=config.VISION_BASE_URL, timeout=180)
        resp = client.chat.completions.create(
            model=config.VISION_MODEL,
            messages=[{"role": "system", "content": system}, {"role": "user", "content": content}],
            temperature=0,
            max_tokens=8000,
        )
        raw = (resp.choices[0].message.content or "").strip()
    except Exception as exc:
        logger.warning("视觉转写 PDF 失败: %s", exc)
        return None

    if not raw:
        return None
    return {"raw_text": raw, "pages": len(images), "extract_mode": "vision"}
